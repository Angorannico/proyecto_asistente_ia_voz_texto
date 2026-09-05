import os
from pathlib import Path
from docx import Document

# Ruta a la carpeta de plantillas
BASE_DIR = Path(r"c:\Users\angel\OneDrive\Desktop\all_projects\Asistente_psicologia")
TEMPLATES_DIR = BASE_DIR / "templates"
ORIGINAL = TEMPLATES_DIR / "plantilla_historia_clinica.docx"
NUEVA = TEMPLATES_DIR / "plantilla_jinja.docx"

def preparar_plantilla():
    if not ORIGINAL.exists():
        print(f"No se encontró el archivo original: {ORIGINAL}")
        return

    doc = Document(ORIGINAL)

    # Reemplazar tabla 1 (Datos del Paciente)
    try:
        tabla_paciente = doc.tables[0]
        # Reemplazar por búsqueda de texto
        mapeo = {
            "NOMBRE DEL PACIENTE": "{{ informacion_paciente.nombre }}",
            "APELLIDOS": "{{ informacion_paciente.apellidos }}",
            "DOCUMENTO DE IDENTIDAD": "{{ informacion_paciente.dni }}",
            "DIRECCIÓN": "{{ informacion_paciente.direccion }}",
            "CORREO ELECTRÓNICO": "{{ informacion_paciente.correo }}",
            "CELULAR": "{{ informacion_paciente.celular }}",
            "FECHA DE NACIMIENTO": "{{ informacion_paciente.fecha_nacimiento }}",
            "CIUDAD": "{{ informacion_paciente.ciudad }}"
        }
        
        # En la primera tabla, las celdas de la derecha de los títulos tienen las "X"
        for row in tabla_paciente.rows:
            for i, cell in enumerate(row.cells):
                texto_celda = cell.text.strip().upper()
                if texto_celda in mapeo and i + 1 < len(row.cells):
                    row.cells[i+1].text = mapeo[texto_celda]
    except Exception as e:
        print(f"Error procesando tabla de paciente: {e}")

    # Lista de reemplazos directos para las tablas de 1 celda (Lorem Ipsum)
    reemplazos = [
        # Problema actual
        ("▶︎ Inicio del malestar", "{{ problema_actual.inicio_malestar }}"),
        ("▶︎ Evolución", "{{ problema_actual.evolucion }}"),
        ("▶︎ Factores desencadenantes", "{{ problema_actual.factores_desencadenantes }}"),
        ("▶︎ Impacto", "{{ problema_actual.impacto_vida_diaria }}"),
        ("▶︎ Síntomas físicos", "{{ problema_actual.sintomas_fisicos }}"),
        
        # Historia Personal
        ("▶︎ Desarrollo", "{{ historia_personal.desarrollo_temprano }}"),
        ("▶︎ Educación", "{{ historia_personal.educacion }}"),
        ("▶︎ Vida laboral", "{{ historia_personal.historia_laboral }}"),
        ("▶︎ Relaciones interpersonales", "{{ historia_personal.relaciones_interpersonales }}"),
        ("▶︎ Historia familiar", "{{ historia_personal.historia_familiar }}"),
        
        # Contexto actual
        ("▶︎ Estructura familiar", "{{ contexto_actual.estructura_familiar }}"),
        ("▶︎ Situación laboral", "{{ contexto_actual.situacion_laboral_academica }}"),
        ("▶︎ Hábitos", "{{ contexto_actual.habitos_salud_estilo_vida }}"),
        ("▶︎ Apoyo social", "{{ contexto_actual.red_apoyo_social }}"),
        
        # Funcionamiento
        ("▶︎ Regulación", "{{ funcionamiento_psicologico.regulacion_emocional }}"),
        ("▶︎ Patrones", "{{ funcionamiento_psicologico.patrones_pensamiento }}"),
        ("▶︎ Comportamiento", "{{ funcionamiento_psicologico.comportamiento }}"),
        ("▶︎ Autoestima", "{{ funcionamiento_psicologico.autoestima_autoimagen }}"),
        
        # Recursos
        ("▶︎ Habilidades de afrontamiento", "{{ recursos_fortalezas.habilidades_afrontamiento }}"),
        ("▶︎ Intereses", "{{ recursos_fortalezas.intereses_pasatiempos }}"),
        ("▶︎ Valores", "{{ recursos_fortalezas.valores_creencias }}"),
        
        # Riesgos
        ("▶︎ Ideación", "{{ evaluacion_riesgo.ideacion_suicida }}"),
        ("▶︎ Riesgo de autolesión", "{{ evaluacion_riesgo.riesgo_autolesion_dano_terceros }}"),
        ("▶︎ Consumo", "{{ evaluacion_riesgo.consumo_sustancias_riesgo }}"),
        
        # Hipótesis
        ("▶︎ Impresión", "{{ hipotesis_clinica.impresion_diagnostica }}"),
        ("▶︎ Formulación", "{{ hipotesis_clinica.formulacion_caso }}"),
        
        # Objetivos
        ("▶︎ Metas", "{{ objetivos_terapeuticos.metas_corto_plazo }}"),
        ("▶︎ Largo", "{{ objetivos_terapeuticos.metas_largo_plazo }}"),
        
        # Plan
        ("▶︎ Enfoque", "{{ plan_intervencion.enfoque_terapeutico }}"),
        ("▶︎ Frecuencia", "{{ plan_intervencion.frecuencia_sesiones }}"),
        ("▶︎ Técnicas", "{{ plan_intervencion.tecnicas_especificas }}"),
        
        # Observaciones
        ("▶︎ Comportamiento durante", "{{ observaciones_clinicas.comportamiento_durante_sesion }}"),
        ("▶︎ Apariencia", "{{ observaciones_clinicas.apariencia_actitud }}"),
    ]

    # Iterar por todos los párrafos y buscar la sección antes de la tabla para saber qué inyectar
    # Como el Lorem Ipsum está dentro de tablas de 1 celda, iteraremos las tablas.
    
    # Índice de tabla actual a modificar (empezamos en 1 porque la 0 es la del paciente)
    tabla_idx = 1
    
    # Recorrer todos los párrafos buscando los títulos
    for p_idx, p in enumerate(doc.paragraphs):
        texto_parrafo = p.text.strip()
        
        # Buscar coincidencias con nuestros reemplazos
        for titulo, tag_jinja in reemplazos:
            if texto_parrafo.startswith(titulo):
                # Si encontramos un título, sabemos que la siguiente tabla le pertenece a este campo.
                if tabla_idx < len(doc.tables):
                    try:
                        # Reemplazar todo el contenido de la tabla (que tiene lorem ipsum) con la variable
                        tabla = doc.tables[tabla_idx]
                        if len(tabla.rows) > 0 and len(tabla.rows[0].cells) > 0:
                            tabla.rows[0].cells[0].text = tag_jinja
                            tabla_idx += 1
                    except Exception as e:
                        print(f"Error reemplazando tabla de {titulo}: {e}")
                break

    # Tratar la firma (última tabla)
    try:
        ultima_tabla = doc.tables[-1]
        ultima_tabla.rows[0].cells[0].text = "{{ profesional.nombre }}\nPROFESIONAL T.P {{ profesional.tarjeta_profesional }}"
    except Exception:
        pass

    doc.save(NUEVA)
    print(f"✅ Plantilla preparada y guardada en {NUEVA}")

if __name__ == "__main__":
    preparar_plantilla()
